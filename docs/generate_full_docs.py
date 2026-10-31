"""
InvoiceIQ Complete Project Documentation Generator
Compiles:
  1. PROJECT_DOCUMENTATION.md (Repository root)
  2. docs/PROJECT_DOCUMENTATION.md (Docs folder)
  3. docs/PROJECT_DOCUMENTATION.docx (Professional Microsoft Word document)
"""

import sys
import os
import re
from pathlib import Path

# Add docs folder to sys.path so it imports doc_data and doc_data_part2
sys.path.insert(0, str(Path(__file__).resolve().parent))

from doc_data import (
    COVER_PAGE, ABSTRACT, INTRODUCTION, PROBLEM_STATEMENT, OBJECTIVES,
    EXISTING_VS_PROPOSED, TECH_STACK, SYSTEM_ARCHITECTURE, MODULE_DESCRIPTION
)
from doc_data_part2 import (
    DETAILED_WORKFLOW, DATABASE_DESIGN, API_DOCUMENTATION, ML_METHODOLOGY,
    INSTALLATION_AND_EXECUTION, TESTING_AND_RESULTS, SAMPLE_IO,
    SECURITY_LIMITATIONS_FUTURE, CONCLUSION, REFERENCES
)

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(__file__).resolve().parent.parent
DOCS_DIR = BASE_DIR / "docs"
DOCS_DIR.mkdir(parents=True, exist_ok=True)

MD_ROOT_PATH = BASE_DIR / "PROJECT_DOCUMENTATION.md"
MD_DOCS_PATH = DOCS_DIR / "PROJECT_DOCUMENTATION.md"
DOCX_PATH = DOCS_DIR / "PROJECT_DOCUMENTATION.docx"


def build_markdown_document() -> str:
    """Stitches together the complete Markdown documentation."""
    md = []
    
    # 1. Cover Page
    md.append(f"# {COVER_PAGE['title']}")
    md.append(f"## {COVER_PAGE['subtitle']}\n")
    md.append("**Academic Project Report & Comprehensive Technical Specification**\n")
    md.append("| Metadata Attribute | Academic / Institutional Details |")
    md.append("|---|---|")
    for k, v in COVER_PAGE["fields"]:
        md.append(f"| **{k}** | {v} |")
    md.append("\n---\n")

    # Table of Contents
    md.append("## Table of Contents")
    toc_items = [
        "1. Cover Page & Project Overview",
        "2. Abstract",
        "3. Introduction",
        "4. Problem Statement",
        "5. Objectives",
        "6. Existing System and Proposed System",
        "7. Technology Stack",
        "8. System Architecture",
        "9. Module Description",
        "10. Detailed Workflow",
        "11. Database Design",
        "12. API Documentation",
        "13. Machine Learning Methodology",
        "14. Installation and Execution",
        "15. Testing and Results",
        "16. Sample Input and Output",
        "17. Security, Limitations, and Future Scope",
        "18. Conclusion",
        "19. References",
    ]
    for item in toc_items:
        md.append(f"- {item}")
    md.append("\n---\n")

    # 2. Abstract
    md.append("## 2. Abstract")
    md.append(ABSTRACT.strip() + "\n\n---\n")

    # 3. Introduction
    md.append("## 3. Introduction")
    md.append(INTRODUCTION.strip() + "\n\n---\n")

    # 4. Problem Statement
    md.append("## 4. Problem Statement")
    md.append(PROBLEM_STATEMENT.strip() + "\n\n---\n")

    # 5. Objectives
    md.append("## 5. Objectives")
    md.append(OBJECTIVES.strip() + "\n\n---\n")

    # 6. Existing vs Proposed
    md.append("## 6. Existing System and Proposed System")
    md.append(EXISTING_VS_PROPOSED.strip() + "\n\n---\n")

    # 7. Tech Stack
    md.append("## 7. Technology Stack")
    md.append(TECH_STACK.strip() + "\n\n---\n")

    # 8. Architecture
    md.append("## 8. System Architecture")
    md.append(SYSTEM_ARCHITECTURE.strip() + "\n\n---\n")

    # 9. Modules
    md.append("## 9. Module Description")
    md.append(MODULE_DESCRIPTION.strip() + "\n\n---\n")

    # 10. Workflow
    md.append("## 10. Detailed Workflow")
    md.append(DETAILED_WORKFLOW.strip() + "\n\n---\n")

    # 11. Database Design
    md.append("## 11. Database Design")
    md.append(DATABASE_DESIGN.strip() + "\n\n---\n")

    # 12. API Documentation
    md.append("## 12. API Documentation")
    md.append(API_DOCUMENTATION.strip() + "\n\n---\n")

    # 13. ML Methodology
    md.append("## 13. Machine Learning Methodology")
    md.append(ML_METHODOLOGY.strip() + "\n\n---\n")

    # 14. Installation & Execution
    md.append("## 14. Installation and Execution")
    md.append(INSTALLATION_AND_EXECUTION.strip() + "\n\n---\n")

    # 15. Testing and Results
    md.append("## 15. Testing and Results")
    md.append(TESTING_AND_RESULTS.strip() + "\n\n---\n")

    # 16. Sample Input & Output
    md.append("## 16. Sample Input and Output")
    md.append(SAMPLE_IO.strip() + "\n\n---\n")

    # 17. Security, Limitations, Future Scope
    md.append("## 17. Security, Limitations, and Future Scope")
    md.append(SECURITY_LIMITATIONS_FUTURE.strip() + "\n\n---\n")

    # 18. Conclusion
    md.append("## 18. Conclusion")
    md.append(CONCLUSION.strip() + "\n\n---\n")

    # 19. References
    md.append("## 19. References")
    md.append(REFERENCES.strip() + "\n")

    return "\n".join(md)


def set_cell_background(cell, hex_color: str):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)


def set_table_borders(table, color="CBD5E1"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)


def add_callout(doc: Document, text: str, title: str = "NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0FDFA")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only (Teal)
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="0D9488"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(13, 148, 136)
    r_title.font.name = "Arial"
    r_title.font.size = Pt(10)
    
    r_text = p.add_run(text)
    r_text.font.name = "Arial"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(51, 65, 85)
    
    # Space after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)


def add_code_block(doc: Document, code_text: str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)


def format_markdown_table_to_docx(doc: Document, md_lines: list):
    """Converts a Markdown table into a styled Word table."""
    raw_rows = [line.strip().strip("|").split("|") for line in md_lines if line.strip() and not re.match(r"^\|?[\s\-:|]+\|?$", line)]
    if not raw_rows:
        return
    
    headers = [c.strip() for c in raw_rows[0]]
    data = [[c.strip() for c in r] for r in raw_rows[1:]]
    
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    # Header row
    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text.replace("**", ""))
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data rows
    for row_idx, row_data in enumerate(data):
        bg = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
        for col_idx in range(len(headers)):
            cell = table.cell(row_idx + 1, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            text_val = row_data[col_idx] if col_idx < len(row_data) else ""
            
            # Handle inline markdown bold/code
            parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text_val)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                elif part.startswith("`") and part.endswith("`"):
                    r = p.add_run(part[1:-1])
                    r.font.name = "Consolas"
                    r.font.size = Pt(8.5)
                else:
                    r = p.add_run(part)
                r.font.name = "Arial"
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(51, 65, 85)
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(8)


def render_content_to_docx(doc: Document, text: str):
    """Parses markdown text into headings, paragraphs, bullet points, code blocks, tables."""
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # Empty line
        if not stripped:
            i += 1
            continue
            
        # Code block
        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # Skip closing ```
            add_code_block(doc, "\n".join(code_lines))
            continue
            
        # Markdown table
        if stripped.startswith("|") and stripped.endswith("|"):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                tbl_lines.append(lines[i].strip())
                i += 1
            format_markdown_table_to_docx(doc, tbl_lines)
            continue
            
        # Heading 3 (###)
        if stripped.startswith("### "):
            h = doc.add_paragraph(style="Heading 3")
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(stripped[4:])
            r.bold = True
            r.font.name = "Arial"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(15, 23, 42) # Slate-900
            i += 1
            continue

        # Heading 2 (##)
        if stripped.startswith("## "):
            h = doc.add_paragraph(style="Heading 2")
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(stripped[3:])
            r.bold = True
            r.font.name = "Arial"
            r.font.size = Pt(14)
            r.font.color.rgb = RGBColor(15, 118, 110) # Teal-700
            i += 1
            continue

        # Heading 1 (#)
        if stripped.startswith("# "):
            h = doc.add_paragraph(style="Heading 1")
            h.paragraph_format.space_before = Pt(22)
            h.paragraph_format.space_after = Pt(8)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(stripped[2:])
            r.bold = True
            r.font.name = "Arial"
            r.font.size = Pt(18)
            r.font.color.rgb = RGBColor(30, 41, 59) # Slate-800
            i += 1
            continue
            
        # Blockquote / Callout (> ...)
        if stripped.startswith("> "):
            callout_text = stripped[2:]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("> "):
                callout_text += " " + lines[i].strip()[2:]
                i += 1
            add_callout(doc, callout_text, title="NOTE / AUDIT NOTICE")
            continue
            
        # Bullet list (- or *)
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            bullet_text = stripped[2:]
            parts = re.split(r"(\*\*.*?\*\*|`.*?`)", bullet_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                elif part.startswith("`") and part.endswith("`"):
                    r = p.add_run(part[1:-1])
                    r.font.name = "Consolas"
                    r.font.size = Pt(9.5)
                else:
                    r = p.add_run(part)
                r.font.name = "Arial"
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Numbered list (1. 2. etc)
        m_num = re.match(r"^(\d+\.)\s+(.*)$", stripped)
        if m_num:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            num_text = m_num.group(2)
            parts = re.split(r"(\*\*.*?\*\*|`.*?`)", num_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                elif part.startswith("`") and part.endswith("`"):
                    r = p.add_run(part[1:-1])
                    r.font.name = "Consolas"
                    r.font.size = Pt(9.5)
                else:
                    r = p.add_run(part)
                r.font.name = "Arial"
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Standard paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        parts = re.split(r"(\*\*.*?\*\*|`.*?`)", stripped)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("`") and part.endswith("`"):
                r = p.add_run(part[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)
            else:
                r = p.add_run(part)
            r.font.name = "Arial"
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(30, 41, 59)
            
        i += 1


def build_docx_document(output_path: Path):
    """Compiles the complete professional Microsoft Word document."""
    doc = Document()
    
    # 1. Page Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header & Footer setup
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("InvoiceIQ — Comprehensive Project Documentation")
        r_hdr.font.name = "Arial"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(148, 163, 184)
        
        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ftr = p_ftr.add_run("Automated Invoice Processing & Ledger Reconciliation System | Confidential")
        r_ftr.font.name = "Arial"
        r_ftr.font.size = Pt(8.5)
        r_ftr.font.color.rgb = RGBColor(148, 163, 184)

    # 2. Cover Page
    # Top spacing
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(40)
    
    # Main Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("InvoiceIQ")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(36)
    r_title.font.color.rgb = RGBColor(15, 118, 110) # Teal-700
    p_title.paragraph_format.space_after = Pt(4)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("AI-Powered Invoice Processing and Ledger Reconciliation System")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(16)
    r_sub.font.color.rgb = RGBColor(71, 85, 105) # Slate-600
    p_sub.paragraph_format.space_after = Pt(20)

    # Decorative Line
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_line = p_line.add_run("—" * 35)
    r_line.font.color.rgb = RGBColor(203, 213, 225)
    p_line.paragraph_format.space_after = Pt(30)

    # Submission Tagline
    p_tag = doc.add_paragraph()
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tag = p_tag.add_run("A Project Report Submitted in Partial Fulfillment of Academic Requirements")
    r_tag.italic = True
    r_tag.font.name = "Arial"
    r_tag.font.size = Pt(11)
    r_tag.font.color.rgb = RGBColor(100, 116, 139)
    p_tag.paragraph_format.space_after = Pt(40)

    # Metadata Table
    tbl = doc.add_table(rows=len(COVER_PAGE["fields"]), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="E2E8F0")
    for idx, (label, val) in enumerate(COVER_PAGE["fields"]):
        c_lbl = tbl.cell(idx, 0)
        c_val = tbl.cell(idx, 1)
        set_cell_background(c_lbl, "F8FAFC")
        set_cell_background(c_val, "FFFFFF")
        set_cell_margins(c_lbl, top=80, bottom=80, left=140, right=140)
        set_cell_margins(c_val, top=80, bottom=80, left=140, right=140)
        
        p0 = c_lbl.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.name = "Arial"
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(51, 65, 85)
        
        p1 = c_val.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(15, 23, 42)
        
    doc.add_page_break()

    # 3. Table of Contents
    p_toc_title = doc.add_paragraph()
    r_tt = p_toc_title.add_run("Table of Contents")
    r_tt.bold = True
    r_tt.font.name = "Arial"
    r_tt.font.size = Pt(20)
    r_tt.font.color.rgb = RGBColor(15, 118, 110)
    p_toc_title.paragraph_format.space_after = Pt(14)
    
    toc_data = [
        ("1", "Cover Page & Academic Credentials", "1"),
        ("2", "Abstract", "2"),
        ("3", "Introduction & Industry Background", "3"),
        ("4", "Problem Statement & Technical Challenges", "5"),
        ("5", "Project Objectives", "6"),
        ("6", "Existing System vs. Proposed InvoiceIQ Architecture", "7"),
        ("7", "Technology Stack Verification", "9"),
        ("8", "System Architecture & End-to-End Data Flow", "11"),
        ("9", "Comprehensive Module Description (11 Modules)", "13"),
        ("10", "Detailed Execution Workflow & Data Provenance", "19"),
        ("11", "Database Design & Relational Schema (ER Model)", "21"),
        ("12", "Complete REST API Documentation", "26"),
        ("13", "Machine Learning Methodology (Isolation Forest)", "32"),
        ("14", "Installation, Configuration & Execution Guide", "36"),
        ("15", "Testing Architecture, Automated Suites & Results", "39"),
        ("16", "Sample Input and Output (Fictional Verification)", "42"),
        ("17", "Security Posture, Current Limitations & Future Scope", "44"),
        ("18", "Conclusion", "47"),
        ("19", "References & Technical Bibliography", "48"),
    ]
    
    tbl_toc = doc.add_table(rows=len(toc_data) + 1, cols=3)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_toc, color="E2E8F0")
    
    # TOC Header
    for c_i, h_txt in enumerate(["#", "Section Name", "Page"]):
        cell = tbl_toc.cell(0, c_i)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h_txt)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_i, (sec_num, sec_name, sec_pg) in enumerate(toc_data):
        bg = "FFFFFF" if r_i % 2 == 0 else "F8FAFC"
        for c_i, val in enumerate([sec_num, sec_name, sec_pg]):
            cell = tbl_toc.cell(r_i + 1, c_i)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=120, right=120)
            p = cell.paragraphs[0]
            if c_i == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

    # Sections Content Rendering
    sections_map = [
        ("2. Abstract", ABSTRACT),
        ("3. Introduction", INTRODUCTION),
        ("4. Problem Statement", PROBLEM_STATEMENT),
        ("5. Objectives", OBJECTIVES),
        ("6. Existing System and Proposed System", EXISTING_VS_PROPOSED),
        ("7. Technology Stack", TECH_STACK),
        ("8. System Architecture", SYSTEM_ARCHITECTURE),
        ("9. Module Description", MODULE_DESCRIPTION),
        ("10. Detailed Workflow", DETAILED_WORKFLOW),
        ("11. Database Design", DATABASE_DESIGN),
        ("12. API Documentation", API_DOCUMENTATION),
        ("13. Machine Learning Methodology", ML_METHODOLOGY),
        ("14. Installation and Execution", INSTALLATION_AND_EXECUTION),
        ("15. Testing and Results", TESTING_AND_RESULTS),
        ("16. Sample Input and Output", SAMPLE_IO),
        ("17. Security, Limitations, and Future Scope", SECURITY_LIMITATIONS_FUTURE),
        ("18. Conclusion", CONCLUSION),
        ("19. References", REFERENCES),
    ]

    for title, content in sections_map:
        # Section Heading
        h = doc.add_paragraph(style="Heading 1")
        h.paragraph_format.space_before = Pt(20)
        h.paragraph_format.space_after = Pt(8)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(title)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(15, 118, 110) # Teal-700
        
        # Render markdown content into styled Word paragraphs/tables
        render_content_to_docx(doc, content)
        doc.add_page_break()

    # Save Word document
    doc.save(str(output_path))
    print(f"Generated Word document: {output_path}")


def main():
    print("=== Generating Markdown Documentation ===")
    md_content = build_markdown_document()
    
    # Write to root directory
    with open(MD_ROOT_PATH, "w", encoding="utf-8") as f:
        f.write(md_content)
    print(f"Wrote: {MD_ROOT_PATH}")

    # Write copy to docs directory
    with open(MD_DOCS_PATH, "w", encoding="utf-8") as f:
        f.write(md_content)
    print(f"Wrote: {MD_DOCS_PATH}")

    print("=== Generating Word (.docx) Documentation ===")
    build_docx_document(DOCX_PATH)

    print("\nDocumentation generation complete!")
    print(f"1. Markdown (Root): {MD_ROOT_PATH}")
    print(f"2. Markdown (Docs): {MD_DOCS_PATH}")
    print(f"3. Word (.docx):    {DOCX_PATH}")


if __name__ == "__main__":
    main()
