"""
InvoiceIQ Complete Academic Project Report Generator
Builds 'InvoiceIQ_Project_Report.docx' in the project root folder.
Modeled after academic B.Tech project reports and formatted with professional typography.
"""

import os
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DOCX_PATH = BASE_DIR / "InvoiceIQ_Project_Report.docx"

# Color Palette (Academic Navy / Teal / Slate)
NAVY_PRIMARY = RGBColor(30, 58, 138)     # #1E3A8A
TEAL_SECONDARY = RGBColor(15, 118, 110)  # #0F766E
SLATE_DARK = RGBColor(30, 41, 59)        # #1E293B
CHARCOAL_BODY = RGBColor(51, 65, 85)     # #334155
MUTED_GRAY = RGBColor(100, 116, 139)     # #64748B
BORDER_COLOR_HEX = "CBD5E1"
BG_LIGHT_HEX = "F8FAFC"
BG_HEADER_HEX = "1E3A8A"
BG_CALLOUT_HEX = "F1F5F9"


def set_cell_background(cell, fill_hex: str):
    """Sets background shading of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Sets inner padding for a table cell in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """Configures individual cell borders."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    borders = {"top": top, "bottom": bottom, "left": left, "right": right}
    for side, border in borders.items():
        if border:
            val, sz, col = border
            borders_xml += f'<{val} {nsdecls("w")} w:val="single" w:sz="{sz}" w:space="0" w:color="{col}"/>'
        else:
            borders_xml += f'<{val} {nsdecls("w")} w:val="none"/>' if False else ""
    borders_xml += '</w:tcBorders>'
    tc_pr.append(parse_xml(borders_xml))


def set_table_borders(table, color="CBD5E1", sz="4"):
    """Applies clean, subtle outer and inner grid borders to a table."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)


def add_callout(doc: Document, text: str, title: str = "KEY ARCHITECTURAL TENET", alert_type: str = "info"):
    """Adds a stylish academic callout box with a colored left accent border."""
    border_color = "1E3A8A" if alert_type == "info" else ("0F766E" if alert_type == "success" else "DC2626")
    bg_color = "F0F9FF" if alert_type == "info" else ("F0FDF4" if alert_type == "success" else "FEF2F2")

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)

    # Thick left accent border
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(tc_borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"[{title}]\n")
    run_t.bold = True
    run_t.font.name = "Arial"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = NAVY_PRIMARY if alert_type == "info" else (TEAL_SECONDARY if alert_type == "success" else RGBColor(220, 38, 38))

    run_b = p.add_run(text)
    run_b.font.name = "Arial"
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = CHARCOAL_BODY

    # Spacer after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)


def add_code_block(doc: Document, code_text: str):
    """Renders formatted monospace code block inside a shaded container."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=140)

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="94A3B8"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(tc_borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(15, 23, 42)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)


def add_figure_placeholder(doc: Document, figure_num: int, title: str, description: str, ascii_mockup: str):
    """Renders a styled screenshot placeholder container with UI layout mockup and formal caption."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="12" w:space="0" w:color="CBD5E1"/>'
        f'  <w:left w:val="single" w:sz="12" w:space="0" w:color="CBD5E1"/>'
        f'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="CBD5E1"/>'
        f'  <w:right w:val="single" w:sz="12" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(tc_borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_hdr = p.add_run(f"APPLICATION SCREEN MOCKUP / WIREFRAME: {title.upper()}\n")
    r_hdr.bold = True
    r_hdr.font.name = "Arial"
    r_hdr.font.size = Pt(9.5)
    r_hdr.font.color.rgb = NAVY_PRIMARY

    # Add ascii mockup
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(6)
    p2.paragraph_format.line_spacing = 1.0
    r_mock = p2.add_run(ascii_mockup.strip())
    r_mock.font.name = "Consolas"
    r_mock.font.size = Pt(8.0)
    r_mock.font.color.rgb = RGBColor(30, 41, 59)

    # Caption paragraph outside table
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(12)
    cap.paragraph_format.keep_with_next = True
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c1 = cap.add_run(f"Figure {figure_num}: {title}. ")
    r_c1.bold = True
    r_c1.font.name = "Arial"
    r_c1.font.size = Pt(9.5)
    r_c1.font.color.rgb = SLATE_DARK

    r_c2 = cap.add_run(description)
    r_c2.font.name = "Arial"
    r_c2.font.size = Pt(9.0)
    r_c2.italic = True
    r_c2.font.color.rgb = MUTED_GRAY


def add_custom_table(doc: Document, headers: list, rows: list, col_widths: list = None):
    """Creates a beautifully styled, full-width academic data table."""
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    set_table_borders(tbl, color="CBD5E1", sz="4")

    # Header Row
    hdr_row = tbl.rows[0]
    tr_pr = hdr_row._tr.get_or_add_trPr()
    tr_pr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    for idx, heading in enumerate(headers):
        c = hdr_row.cells[idx]
        if col_widths and idx < len(col_widths):
            c.width = Inches(col_widths[idx])
        set_cell_background(c, BG_HEADER_HEX)
        set_cell_margins(c, top=140, bottom=140, left=140, right=140)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(heading)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        row = tbl.rows[r_idx + 1]
        bg = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
        for c_idx, cell_value in enumerate(row_data):
            c = row.cells[c_idx]
            if col_widths and c_idx < len(col_widths):
                c.width = Inches(col_widths[c_idx])
            set_cell_background(c, bg)
            set_cell_margins(c, top=100, bottom=100, left=140, right=140)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(cell_value))
            r.font.name = "Arial"
            r.font.size = Pt(9.0)
            r.font.color.rgb = CHARCOAL_BODY

    # Spacer
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)


def configure_header_footer(doc: Document):
    """Configures running header and footer across document sections."""
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.different_first_page_header_footer = True

    # Running Header (Pages 2+)
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(6)
    hr = hp.add_run("InvoiceIQ: AI-Powered Invoice Processing and Ledger Reconciliation System")
    hr.font.name = "Arial"
    hr.font.size = Pt(8.5)
    hr.font.color.rgb = MUTED_GRAY

    # Running Footer (Pages 2+)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fp.paragraph_format.space_before = Pt(6)
    fr_left = fp.add_run("B.Tech Project Report — Department of Computer Science & Engineering\t\t")
    fr_left.font.name = "Arial"
    fr_left.font.size = Pt(8.5)
    fr_left.font.color.rgb = MUTED_GRAY

    # Add dynamic page number field in footer
    f_run = fp.add_run("Page ")
    f_run.font.name = "Arial"
    f_run.font.size = Pt(8.5)
    f_run.font.color.rgb = MUTED_GRAY

    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    fp._p.append(fldSimple)

    f_run2 = fp.add_run(" of ")
    f_run2.font.name = "Arial"
    f_run2.font.size = Pt(8.5)
    f_run2.font.color.rgb = MUTED_GRAY

    fldSimple2 = OxmlElement('w:fldSimple')
    fldSimple2.set(qn('w:instr'), 'NUMPAGES')
    fp._p.append(fldSimple2)


print("Initialized helper modules for academic project report generation.")
